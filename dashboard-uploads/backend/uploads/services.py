"""
Validation + text extraction for multimodal uploads.

Documents (.pdf .txt .docx) are extracted here.
Images (OCR) and audio (speech-to-text) are delegated to OPTIONAL handlers that you
plug in from your OCR_modul / voice_search code via settings.py:

    UPLOAD_OCR_HANDLER        = "OCR_modul.api.extract_text"      # f(path: str) -> str
    UPLOAD_TRANSCRIBE_HANDLER = "voice_search.api.transcribe"     # f(path: str) -> str

If a handler is not configured the file is still saved (status "stored").
Everything runs locally; nothing is sent to an external service.
"""
import os

from django.conf import settings
from django.utils.module_loading import import_string

from .models import Upload

MAX_CHARS = 50_000  # same cap as the existing Document Assistant

MODALITIES = {
    Upload.Modality.DOCUMENT: {"exts": {".pdf", ".txt", ".docx"}, "max_mb": 25},
    Upload.Modality.IMAGE: {"exts": {".png", ".jpg", ".jpeg", ".webp"}, "max_mb": 15},
    Upload.Modality.AUDIO: {"exts": {".mp3", ".wav", ".m4a", ".ogg"}, "max_mb": 100},
}


def detect_modality(filename):
    ext = os.path.splitext(filename)[1].lower()
    for modality, cfg in MODALITIES.items():
        if ext in cfg["exts"]:
            return modality, ext
    return None, ext


def validate(file):
    """Returns (error_message_or_None, modality_or_None)."""
    modality, ext = detect_modality(file.name)
    if modality is None:
        allowed = ", ".join(sorted(e for c in MODALITIES.values() for e in c["exts"]))
        return f"Unsupported file type '{ext or file.name}'. Allowed: {allowed}", None
    if file.size == 0:
        return "The file is empty.", None
    limit = MODALITIES[modality]["max_mb"]
    if file.size > limit * 1024 * 1024:
        return f"File is larger than the {limit} MB limit for {modality} files.", None
    return None, modality


# ---------- extractors ----------
def _read_pdf(path):
    import fitz  # PyMuPDF

    with fitz.open(path) as doc:
        return "\n".join(page.get_text() for page in doc)


def _read_txt(path):
    with open(path, "r", encoding="utf-8", errors="replace") as fh:
        return fh.read()


def _read_docx(path):
    from docx import Document  # python-docx (already used for report generation)

    d = Document(path)
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for table in d.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def _call_handler(setting_name, path):
    dotted = getattr(settings, setting_name, None)
    if not dotted:
        return None
    return import_string(dotted)(path)


def process(upload):
    """Fills extracted_text / status / note on the Upload (caller saves it)."""
    path = upload.file.path
    ext = os.path.splitext(upload.original_name)[1].lower()
    try:
        if upload.modality == Upload.Modality.DOCUMENT:
            reader = {".pdf": _read_pdf, ".txt": _read_txt, ".docx": _read_docx}[ext]
            text = reader(path)
        elif upload.modality == Upload.Modality.IMAGE:
            text = _call_handler("UPLOAD_OCR_HANDLER", path)
            if text is None:
                upload.status, upload.note = Upload.Status.STORED, "Saved. No OCR handler configured yet."
                return
        else:  # audio
            text = _call_handler("UPLOAD_TRANSCRIBE_HANDLER", path)
            if text is None:
                upload.status, upload.note = Upload.Status.STORED, "Saved. No transcription handler configured yet."
                return

        text = (text or "").strip()
        if not text:
            upload.status, upload.note = Upload.Status.STORED, "Saved, but no text was found in the file."
            return
        if len(text) > MAX_CHARS:
            text = text[:MAX_CHARS]
            upload.note = f"Text truncated to {MAX_CHARS:,} characters."
        upload.extracted_text = text
        upload.status = Upload.Status.PROCESSED
    except Exception as exc:  # keep the file, report the failure
        upload.status = Upload.Status.FAILED
        upload.note = f"Could not process file: {exc}"
